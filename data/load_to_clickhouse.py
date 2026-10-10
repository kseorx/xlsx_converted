"""Load selected Excel budget forms into ClickHouse."""

from __future__ import annotations

import argparse
import hashlib
import re
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

import config


INSERT_COLUMNS = [
    "id", "pp_number", "article_code", "article_name", "ekbs_code", "metric_unit",
    "value_year_minus2", "value_year_minus1", "value_current_year", "value_jan",
    "value_feb", "value_mar", "value_q1", "value_apr", "value_may", "value_jun",
    "value_q2", "value_jul", "value_aug", "value_sep", "value_q3", "value_oct",
    "value_nov", "value_dec", "value_q4", "value_6m", "value_9m", "date_year",
    "form_type", "organization_name", "info_type", "source_id", "version_id",
]
SQLITE_COLUMN_TYPES = {
    column: "INTEGER" if column == "id"
    else "REAL" if column.startswith("value_")
    else "TEXT"
    for column in INSERT_COLUMNS
}
SQLITE_CREATE_TABLE_SQL = (
    "CREATE TABLE IF NOT EXISTS budget_reports ("
    + ", ".join(f"{column} {SQLITE_COLUMN_TYPES[column]}" for column in INSERT_COLUMNS)
    + ")"
)

MONTHS = {
    "jan": ("январ", "january"),
    "feb": ("феврал", "february"),
    "mar": ("март", "march"),
    "apr": ("апрел", "april"),
    "may": ("май", "мая", "майский", "may"),
    "jun": ("июн", "june"),
    "jul": ("июл", "july"),
    "aug": ("август", "авг", "august"),
    "sep": ("сентябр", "сент", "september"),
    "oct": ("октябр", "окт", "october"),
    "nov": ("ноябр", "ноя", "november"),
    "dec": ("декабр", "дек", "december"),
}


def text(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\xa0", " ").split()).strip()


def normalized(value: Any) -> str:
    return text(value).casefold().replace("ё", "е")


def as_string(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    result = text(value)
    return result or None


def as_number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    result = text(value).replace(" ", "")
    if not result or result.casefold() in {"-", "x", "х", "n/a"}:
        return None
    if "," in result and "." not in result:
        result = result.replace(",", ".")
    try:
        return float(result)
    except ValueError:
        return None


def resolve_form(sheet_name: str) -> str | None:
    normalized_name = normalized(sheet_name)
    for pattern, canonical in config.FORM_PATTERNS.items():
        prefix, number = normalized(pattern).split(".", 1)
        match = re.match(
            r"^\s*" + re.escape(prefix) + r"\s*\.\s*" + re.escape(number) + r"(?!\d)",
            normalized_name,
        )
        if match:
            return canonical
    return None


def file_metadata(path: Path) -> dict[str, Any] | None:
    override = config.FILE_METADATA.get(path.name)
    if override:
        organization = text(override.get("organization"))
        year = override.get("year")
        report_type = normalized(override.get("type"))
        if organization and isinstance(year, int) and report_type in {"план", "факт"}:
            return {"organization": organization, "year": year, "type": report_type}
        raise ValueError(f"Некорректный FILE_METADATA для {path.name}")

    stem = path.stem
    year_match = re.search(r"(20\d{2})", stem)
    if not year_match:
        return None
    filename_text = normalized(stem)
    if "факт" in filename_text:
        report_type = "факт"
    elif "план" in filename_text:
        report_type = "план"
    else:
        return None
    organization = stem[:year_match.start()].strip(" _-.")
    if not organization:
        return None
    return {
        "organization": organization,
        "year": int(year_match.group(1)),
        "type": report_type,
    }


def find_header(rows: list[tuple[Any, ...]]) -> tuple[int, dict[str, int]] | None:
    for row_index, row in enumerate(rows):
        columns: dict[str, int] = {}
        for column_index, value in enumerate(row):
            label = normalized(value)
            if not label:
                continue
            if "№" in label and ("п/п" in label or "п.п" in label):
                columns.setdefault("seq_num", column_index)
            if "код статьи" in label:
                columns.setdefault("article_code", column_index)
            if "наименование статьи" in label or "статья формы" in label:
                columns.setdefault("article_name", column_index)
            if "екбс" in label or "ас фзд" in label:
                columns.setdefault("ekbs_code", column_index)
            if "ед. изм" in label or "единицы измерения" in label:
                columns.setdefault("unit", column_index)
        if "seq_num" in columns and "article_name" in columns:
            return row_index, columns
    return None


def header_labels(rows: list[tuple[Any, ...]], header_index: int) -> list[str]:
    start = max(0, header_index - 1)
    stop = min(len(rows), header_index + 3)
    width = max((len(row) for row in rows[start:stop]), default=0)
    return [
        " ".join(text(rows[index][column]) for index in range(start, stop)
                 if column < len(rows[index]) and text(rows[index][column]))
        for column in range(width)
    ]


def _contains_any(value: str, choices: tuple[str, ...] | list[str]) -> bool:
    return any(choice in value for choice in choices)


def metric_columns(
    labels: list[str], metadata_columns: dict[str, int], year: int, report_type: str
) -> dict[str, int]:
    first_value_column = max(metadata_columns.values(), default=-1) + 1
    candidates = {
        index: normalized(label)
        for index, label in enumerate(labels)
        if index >= first_value_column
    }
    month_hits = {
        month: [index for index, label in candidates.items()
                if str(year) in label and _contains_any(label, tokens)]
        for month, tokens in MONTHS.items()
    }
    metrics: dict[str, int] = {}

    def matching_year(target_year: int, predicate: Any) -> list[int]:
        return [index for index, label in candidates.items()
                if str(target_year) in label and predicate(label)]

    def pick(columns: list[int]) -> int | None:
        return columns[0] if columns else None

    current_tokens = ("факт", "исполн") if report_type == "факт" else ("план", "бюджет")
    for offset, field in ((2, "value_year_minus2"), (1, "value_year_minus1")):
        column = pick(matching_year(year - offset, lambda label: True))
        if column is not None:
            metrics[field] = column

    current = matching_year(year, lambda label: _contains_any(label, current_tokens))
    if not current:
        current = matching_year(year, lambda label: True)
    column = pick(current)
    if column is not None:
        metrics["value_current_year"] = column

    for month, tokens in MONTHS.items():
        columns = month_hits[month]
        selected = [index for index in columns if _contains_any(candidates[index], current_tokens)]
        column = pick(selected or columns)
        if column is not None:
            metrics[f"value_{month}"] = column

    quarter_tokens = {
        "value_q1": ("1 кв", "1 кварт"),
        "value_q2": ("2 кв", "2 кварт"),
        "value_q3": ("3 кв", "3 кварт"),
        "value_q4": ("4 кв", "4 кварт"),
        "value_6m": ("6 мес", "полугод"),
        "value_9m": ("9 мес",),
    }
    for field, tokens in quarter_tokens.items():
        column = pick([index for index, label in candidates.items()
                       if _contains_any(label, tokens)])
        if column is not None:
            metrics[field] = column
    return metrics


def stable_report_id(file_name: str, sheet_name: str, row_num: int) -> int:
    identity = f"{file_name}\0{sheet_name}\0{row_num}".encode("utf-8")
    return int.from_bytes(hashlib.sha256(identity).digest()[:8], "big") & ((1 << 63) - 1)


def file_version_id(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def extract_sheet(
    rows: list[tuple[Any, ...]], file_name: str, sheet_name: str, form: str,
    metadata: dict[str, Any], version_id: str,
) -> list[tuple[Any, ...]]:
    header = find_header(rows)
    if header is None:
        print(f"  Пропуск листа без строки заголовков: {sheet_name}", file=sys.stderr)
        return []
    header_index, columns = header
    labels = header_labels(rows, header_index)
    value_columns = metric_columns(labels, columns, metadata["year"], metadata["type"])
    data: list[tuple[Any, ...]] = []

    for row_index in range(header_index + 1, len(rows)):
        row = rows[row_index]
        seq = row[columns["seq_num"]] if columns["seq_num"] < len(row) else None
        name_index = columns["article_name"]
        article_name = row[name_index] if name_index < len(row) else None
        code_index = columns.get("article_code", -1)
        article_code = row[code_index] if 0 <= code_index < len(row) else None
        if "№" in normalized(seq) or normalized(seq) in {"0"} and normalized(article_name) == "2":
            continue
        if seq is None and article_code is None and article_name is None:
            continue
        if not text(article_name):
            continue

        record: dict[str, Any] = {column: None for column in INSERT_COLUMNS}
        record.update({
            "id": stable_report_id(file_name, sheet_name, row_index + 1),
            "pp_number": as_string(seq),
            "article_code": as_string(article_code),
            "article_name": text(article_name),
            "ekbs_code": as_string(row[columns["ekbs_code"]])
            if columns.get("ekbs_code", len(row)) < len(row) else None,
            "metric_unit": as_string(row[columns["unit"]])
            if columns.get("unit", len(row)) < len(row) else None,
            "date_year": str(metadata["year"]),
            "form_type": form,
            "organization_name": metadata["organization"],
            "info_type": metadata["type"],
            "source_id": file_name,
            "version_id": version_id,
        })
        for field, column in value_columns.items():
            if column < len(row):
                record[field] = as_number(row[column])
        data.append(tuple(record[column] for column in INSERT_COLUMNS))
    return data


def process_file(
    path: Path, client: Any | None, dry_run: bool, adapter: str
) -> tuple[int, int]:
    metadata = file_metadata(path)
    if metadata is None:
        print(f"Пропуск файла: не удалось определить организацию, год и тип: {path.name}", file=sys.stderr)
        return 0, 0

    workbook = load_workbook(path, read_only=True, data_only=True)
    file_rows: list[tuple[Any, ...]] = []
    matched_sheets = 0
    version_id = file_version_id(path)
    try:
        for sheet in workbook.worksheets:
            form = resolve_form(sheet.title)
            if form is None:
                continue
            matched_sheets += 1
            sheet_rows = list(sheet.iter_rows(values_only=True))
            file_rows.extend(
                extract_sheet(sheet_rows, path.name, sheet.title, form, metadata, version_id)
            )
    finally:
        workbook.close()

    if matched_sheets == 0:
        print(f"Нет целевых форм в файле: {path.name}", file=sys.stderr)
    if dry_run:
        print(f"{path.name}: {len(file_rows)} строк, формы: {matched_sheets}")
    elif file_rows and client is not None:
        if adapter == "sqlite":
            placeholders = ", ".join("?" for _ in INSERT_COLUMNS)
            statement = (
                f"INSERT INTO budget_reports ({', '.join(INSERT_COLUMNS)}) "
                f"VALUES ({placeholders})"
            )
        else:
            statement = f"INSERT INTO budget_reports ({', '.join(INSERT_COLUMNS)}) VALUES"
        for start in range(0, len(file_rows), config.BATCH_SIZE):
            batch = file_rows[start:start + config.BATCH_SIZE]
            if adapter == "sqlite":
                client.executemany(statement, batch)
            else:
                client.execute(statement, batch)
        print(f"Загружено {len(file_rows)} строк из {path.name}")
    return len(file_rows), matched_sheets


def main() -> int:
    parser = argparse.ArgumentParser(description="Загрузка целевых Excel-форм в SQL-таблицу")
    parser.add_argument("--data-dir", type=Path, default=Path(config.DATA_DIR))
    parser.add_argument("--dry-run", action="store_true", help="Проверить файлы без подключения к ClickHouse")
    parser.add_argument(
        "--sqlite-test",
        nargs="?",
        const=":memory:",
        metavar="DB_PATH",
        help="Загрузить в локальный SQLite вместо ClickHouse; без DB_PATH используется память",
    )
    parser.add_argument(
        "--create-table-only",
        action="store_true",
        help="Создать таблицу ClickHouse без загрузки Excel-строк",
    )
    args = parser.parse_args()
    if args.create_table_only and (args.dry_run or config.DRY_RUN):
        parser.error("--create-table-only нельзя использовать в режиме DRY_RUN")
    if args.sqlite_test is not None and (args.dry_run or config.DRY_RUN or args.create_table_only):
        parser.error("--sqlite-test нельзя совмещать с --dry-run или --create-table-only")
    dry_run = bool(config.DRY_RUN or args.dry_run)

    paths = sorted(
        path for path in args.data_dir.rglob("*")
        if path.is_file() and path.suffix.casefold() in {".xlsx", ".xlsm"}
        and not path.name.startswith("~$")
    )
    if not paths and not args.create_table_only:
        print(f"Excel-файлы не найдены: {args.data_dir}", file=sys.stderr)
        return 1

    client = None
    adapter = "clickhouse"
    if args.sqlite_test is not None:
        import sqlite3

        adapter = "sqlite"
        client = sqlite3.connect(args.sqlite_test)
        client.execute(SQLITE_CREATE_TABLE_SQL)
        database_label = "in-memory" if args.sqlite_test == ":memory:" else args.sqlite_test
        print(f"SQLite test database: {database_label}")
    elif not dry_run:
        try:
            from clickhouse_driver import Client
        except ImportError:
            print("Не установлен clickhouse-driver. Выполните: pip install -r requirements.txt", file=sys.stderr)
            return 1
        client = Client(
            host=config.CH_HOST,
            port=config.CH_PORT,
            user=config.CH_USER,
            password=config.CH_PASSWORD,
            database=config.CH_DATABASE,
            secure=config.CH_SECURE,
        )
        client.execute(config.CREATE_TABLE_SQL)
        if args.create_table_only:
            client.disconnect()
            print("Таблица budget_reports создана или уже существует")
            return 0

    total_rows = 0
    processed_files = 0
    for path in paths:
        try:
            rows_count, _ = process_file(path, client, dry_run, adapter)
            total_rows += rows_count
            processed_files += 1
        except Exception as error:
            print(f"Ошибка при обработке {path.name}: {error}", file=sys.stderr)
            if client is not None:
                client.close() if adapter == "sqlite" else client.disconnect()
            return 1
    if client is not None:
        if adapter == "sqlite":
            client.commit()
            row_count = client.execute("SELECT count(*) FROM budget_reports").fetchone()[0]
            print(f"Строк в SQLite budget_reports: {row_count}")
            for form, count in client.execute(
                "SELECT form_type, count(*) FROM budget_reports GROUP BY form_type ORDER BY form_type"
            ):
                print(f"  {form}: {count}")
            client.close()
        else:
            client.disconnect()
    action = "К проверке" if dry_run else "Всего загружено"
    print(f"{action}: {total_rows} строк; обработано файлов: {processed_files}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())